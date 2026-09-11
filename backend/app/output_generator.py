import os
import json
import uuid
from pathlib import Path
from typing import Dict, Any, Union, Optional
from app.config.settings import settings

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

import pptx
from pptx.util import Inches as PptxInches, Pt as PptxPt
from pptx.dml.color import RGBColor as PptxRGBColor

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def generate_docx(content: str, filename_prefix: str = "Deliverable") -> Path:
    """Generates a professional Word document (.docx)."""
    outputs_dir = settings.OUTPUTS_DIR
    filename = f"{filename_prefix}_{uuid.uuid4().hex[:6]}.docx"
    filepath = outputs_dir / filename

    doc = docx.Document()
    
    # Header styling
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("SOVEREIGN INDUSTRIAL WORKBENCH DELIVERABLE")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(11, 40, 90)

    subtitle_p = doc.add_paragraph()
    sub_run = subtitle_p.add_run("Air-Gapped Confidential Operations Report")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph("-" * 60)

    # Process paragraphs and headers
    lines = str(content).split("\n")
    for line in lines:
        stripped = line.strip()
        if stripped.startswith("# "):
            h = doc.add_heading(stripped[2:], level=1)
            h.runs[0].font.color.rgb = RGBColor(11, 40, 90)
        elif stripped.startswith("## "):
            h = doc.add_heading(stripped[3:], level=2)
            h.runs[0].font.color.rgb = RGBColor(30, 80, 150)
        elif stripped.startswith("### "):
            h = doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            doc.add_paragraph(stripped[2:], style='List Bullet')
        elif stripped:
            doc.add_paragraph(stripped)

    doc.save(str(filepath))
    return filepath

def generate_pptx(slides_data: Union[str, list], filename_prefix: str = "Presentation") -> Path:
    """Generates a PowerPoint presentation (.pptx)."""
    outputs_dir = settings.OUTPUTS_DIR
    filename = f"{filename_prefix}_{uuid.uuid4().hex[:6]}.pptx"
    filepath = outputs_dir / filename

    prs = pptx.Presentation()

    # Parse slides data if string JSON
    if isinstance(slides_data, str):
        try:
            slides_data = json.loads(slides_data)
        except Exception:
            # Fallback text to slides
            slides_data = [{"title": "Industrial Operations Summary", "bullets": slides_data.split("\n")}]

    if not isinstance(slides_data, list):
        slides_data = [{"title": "Report Summary", "bullets": [str(slides_data)]}]

    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = "Sovereign Industrial Report"
    if slide.placeholders[1]:
        slide.placeholders[1].text = "Air-Gapped Automated Deliverable"

    # Content Slides
    bullet_slide_layout = prs.slide_layouts[1]
    for slide_info in slides_data:
        if isinstance(slide_info, dict):
            title_text = slide_info.get("title", "Industrial Insight")
            bullets = slide_info.get("bullets", [])
        else:
            title_text = "Key Takeaway"
            bullets = [str(slide_info)]

        slide = prs.slides.add_slide(bullet_slide_layout)
        slide.shapes.title.text = title_text
        tf = slide.placeholders[1].text_frame
        tf.clear()

        for idx, bullet in enumerate(bullets):
            p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
            p.text = str(bullet)
            p.font.size = PptxPt(16)

    prs.save(str(filepath))
    return filepath

def generate_xlsx(data: Union[str, dict], filename_prefix: str = "DataSheet") -> Path:
    """Generates an Excel spreadsheet (.xlsx)."""
    outputs_dir = settings.OUTPUTS_DIR
    filename = f"{filename_prefix}_{uuid.uuid4().hex[:6]}.xlsx"
    filepath = outputs_dir / filename

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Industrial Audit Data"

    # Styling
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    if isinstance(data, str):
        try:
            data = json.loads(data)
        except Exception:
            data = {"headers": ["Item / Description"], "rows": [[line] for line in data.split("\n") if line.strip()]}

    headers = data.get("headers", ["Column 1", "Column 2"]) if isinstance(data, dict) else ["Value"]
    rows = data.get("rows", []) if isinstance(data, dict) else []

    # Write Headers
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Write Rows
    for row in rows:
        if isinstance(row, list):
            ws.append(row)
        else:
            ws.append([str(row)])

    # Format cells
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.font = cell_font
            cell.border = thin_border

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    wb.save(str(filepath))
    return filepath

def generate_deliverable(task_id: str, format_type: str, content: Any) -> Optional[Dict[str, str]]:
    """
    Main deliverable dispatcher.
    Returns dictionary with download URL and file metadata if deliverable created.
    """
    fmt = str(format_type).strip().lower()
    if fmt not in ["docx", "pptx", "xlsx"]:
        return None

    try:
        if fmt == "docx":
            path = generate_docx(content, filename_prefix=f"Task_{task_id[:6]}")
        elif fmt == "pptx":
            path = generate_pptx(content, filename_prefix=f"Task_{task_id[:6]}")
        elif fmt == "xlsx":
            path = generate_xlsx(content, filename_prefix=f"Task_{task_id[:6]}")
        else:
            return None

        return {
            "format": fmt,
            "filename": path.name,
            "download_url": f"/outputs/{path.name}"
        }
    except Exception as e:
        print(f"Error generating {fmt} deliverable: {e}")
        return None
